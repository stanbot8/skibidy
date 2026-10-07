"""Reproduce descriptive complete-cycle analysis without changing source workbooks.

Requires openpyxl, numpy and scipy. Run from any working directory. Outputs
are data derivatives, not simulation parameter recommendations.
"""
import argparse
import csv
import hashlib
import json
import math
import re
import platform
import subprocess
from collections import Counter
from pathlib import Path
from urllib.request import Request, urlopen

import numpy as np
import openpyxl
import scipy
from scipy import stats

HERE = Path(__file__).resolve().parent


def default_cache_dir():
    try:
        git_dir = subprocess.check_output(
            ['git', 'rev-parse', '--absolute-git-dir'], cwd=HERE, text=True,
            stderr=subprocess.DEVNULL).strip()
    except (OSError, subprocess.CalledProcessError) as error:
        raise RuntimeError('Use --cache-dir for an extraction outside a Git checkout.') from error
    return Path(git_dir) / 'research' / 'tissue'


def verify_source(path, source):
    if path.stat().st_size != source['bytes'] or hashlib.sha256(path.read_bytes()).hexdigest() != source['sha256']:
        raise ValueError(f'Original supplement size/hash mismatch: {path}; retained without modification.')


def source_paths(provenance, cache_dir, download=False):
    paths = {}
    for source in provenance['sources']:
        path = (HERE if source['storage'] == 'tracked' else cache_dir) / source['file']
        if not path.is_file():
            if not download:
                raise FileNotFoundError(f'Missing {path}. Re-run with --download to retrieve the recorded public URL.')
            # Ordinary publisher URL only. Reject changed, truncated or HTML
            # responses before saving; do not overwrite any preserved evidence.
            request = Request(source['download_url'], headers={'User-Agent': 'Skibidy-research-data/1.0'})
            with urlopen(request, timeout=60) as response:
                data = response.read(source['bytes'] + 1)
            if len(data) != source['bytes'] or hashlib.sha256(data).hexdigest() != source['sha256']:
                raise ValueError(f'Publisher response size/hash differs for {source["file"]}; nothing saved.')
            path.parent.mkdir(parents=True, exist_ok=True)
            with path.open('xb') as stream:
                stream.write(data)
        verify_source(path, source)
        paths[source['file']] = path
    return paths


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def extract(study, path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet in workbook:
        rows = list(sheet.iter_rows())
        width = 5 if study == 'Xiao2024' else 7
        offset = 0 if study == 'Xiao2024' else 1
        # Roshan has occasional spacer columns. Resolve each clone from its
        # actual header, rather than assuming every block starts on a stride.
        for col in range(offset, sheet.max_column):
            clone = rows[0][col].value
            if not isinstance(clone, str) or not re.match(r'(Colony|Clone)\s+\d+', clone):
                continue
            duration_col = col + (3 if width == 5 else 5)
            for row in rows[2:]:
                gen = row[col].value
                source_generation = gen
                if isinstance(gen, str) and re.fullmatch(r'G\d+', gen):
                    gen = int(gen[1:])
                if not numeric(gen) and not (isinstance(gen, str) and re.fullmatch(r'(Outer|D)\d+', gen)):
                    continue
                cell = row[duration_col]
                duration = cell.value
                fate = row[duration_col + 1].value
                records.append(dict(study=study, sheet=sheet.title, clone=clone,
                                    generation=int(gen) if numeric(gen) else gen,
                                    source_generation=source_generation, duration_cell=cell.coordinate,
                                    duration_h=duration if numeric(duration) else '',
                                    division_type=fate or '',
                                    positive_recorded_interval=bool(numeric(duration) and duration > 0),
                                    complete_cycle=bool(numeric(duration) and duration > 0 and
                                                        isinstance(fate, str) and re.fullmatch(r'[PDU]{2}', fate)),
                                    first_generation=bool(gen == 1),
                                    fit_eligible=bool(numeric(gen) and gen > 1 and
                                                      sheet.title != 'NFSK expanding (Fig2a)')))
    workbook.close()
    return records


def describe(records):
    selected = [r for r in records if r['complete_cycle'] and r['fit_eligible']]
    x = np.array([r['duration_h'] for r in selected], dtype=float)
    if not len(x):
        return None
    result = dict(n_positive_recorded_intervals=sum(r['positive_recorded_interval'] and
                                                   numeric(r['generation']) and r['generation'] > 1 for r in records),
                  n_complete_cycles=len(x), n_clones_with_complete_cycles=len({r['clone'] for r in selected}),
                  mean_h=float(x.mean()), median_h=float(np.median(x)), sample_sd_h=float(x.std(ddof=1)),
                  observed_cv=float(x.std(ddof=1) / x.mean()), min_h=float(x.min()), max_h=float(x.max()),
                  q25_h=float(np.quantile(x, .25)), q75_h=float(np.quantile(x, .75)),
                  n_over_48h=int((x > 48).sum()), division_types=dict(Counter(r['division_type'] for r in selected)))
    fits = {}
    for name, distribution, options in [('normal', stats.norm, {}),
                                         ('lognormal', stats.lognorm, {'floc': 0}),
                                         ('gamma', stats.gamma, {'floc': 0})]:
        params = distribution.fit(x, **options)
        loglik = float(distribution.logpdf(x, *params).sum())
        # All three families have two estimated parameters; positive families fix loc=0.
        fits[name] = dict(parameters=[float(p) for p in params], log_likelihood=loglik,
                          aic=4 - 2 * loglik)
    result['descriptive_fits'] = fits
    return result


def mouse_inventory(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    inventory = []
    for sheet in workbook:
        if not sheet.title.startswith('H2BGFP_'):
            continue
        rows = list(sheet.iter_rows())
        for row in rows[:10]:
            for cell in row:
                if cell.value != 'Cell Intensity':
                    continue
                col = cell.column - 1
                data = [(r[col].value, r[col - 2].value) for r in rows[cell.row:] if numeric(r[col].value)]
                time_label = next((r[col - 2].value for r in reversed(rows[:cell.row - 1])
                                   if isinstance(r[col - 2].value, str) and 'days' in r[col - 2].value), None)
                inventory.append(dict(sheet=sheet.title, intensity_header=cell.coordinate,
                                      time_label=time_label, n_intensities=len(data),
                                      n_mouse_ids=len({mouse for _, mouse in data if mouse is not None})))
    workbook.close()
    return inventory


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cache-dir', type=Path, help='Local original-supplement cache (default: Git directory/research/tissue).')
    parser.add_argument('--download', action='store_true', help='Fetch missing originals from their recorded public publisher URLs, verifying size and SHA-256.')
    args = parser.parse_args()
    provenance = json.loads((HERE / 'provenance.json').read_text(encoding='utf-8'))
    paths = source_paths(provenance, (args.cache_dir or default_cache_dir()).resolve(), args.download)
    records = extract('Xiao2024', paths['xiao2024_supplement3.xlsx'])
    records += extract('Roshan2016', paths['roshan2016_supplement_table1.xlsx'])
    keys = {(r['study'], r['sheet'], r['duration_cell']) for r in records}
    assert len(keys) == len(records), 'Duplicate source-cell extraction'
    with (HERE / 'human_complete_cycle_records.csv').open('w', newline='', encoding='utf-8') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    summaries = {}
    for study, sheet in sorted({(r['study'], r['sheet']) for r in records}):
        summaries[f'{study}: {sheet}'] = describe([r for r in records if r['study'] == study and r['sheet'] == sheet])
    classified = {}
    for study, sheet in sorted({(r['study'], r['sheet']) for r in records}):
        subset = [r for r in records if r['study'] == study and r['sheet'] == sheet
                  and r['division_type'] in {'PP', 'PD', 'DP', 'DD'}]
        classified[f'{study}: {sheet}'] = describe(subset)
    assert classified['Xiao2024: EXP']['n_complete_cycles'] == 1017
    assert classified['Xiao2024: DIF']['n_complete_cycles'] == 174
    summary = dict(selection='Positive numeric recorded delta-T, numeric local generation > 1, and two daughter-fate labels from P/D/U to establish division. Single D/U, blank and other labels retained but not treated as confirmed completed cycles. Classified subset restricts source fates to PP/PD/DP/DD. NFSK expanding Fig2a has reset subclone generations and unresolved cell identity, so is excluded from fits.',
                   sd_definition='Sample SD, ddof=1; CV is SD / arithmetic mean of observed complete cycles.',
                   fit_definition='Exploratory marginal MLE; normal(mu,sigma), lognormal(shape,loc=0,scale), gamma(shape,loc=0,scale); AIC=4-2logL.',
                   limitations='Complete cycles only; clustered lineages and censored/untracked cells. AIC assumes independence and is descriptive. No latent propensity CV is identified.',
                   neonatal_reconciliation='The article reports 2127 post-plating complete cycles in 81 colonies. Table 1 has subclone generation resets and Outer labels; this extraction does not reconstruct the complete published neonatal cohort and does not fit its expanding sheet. Do not remove repeated times as duplicates without biological cell identity.',
                   runtime=dict(python=platform.python_version(), numpy=np.__version__, scipy=scipy.__version__, openpyxl=openpyxl.__version__),
                   source_hashes={name: hashlib.sha256(paths[name].read_bytes()).hexdigest() for name in sorted(paths)},
                   human_groups=summaries, classified_human_groups=classified,
                   mouse_intensity_inventory=mouse_inventory(paths['piedrafita2020_supplement_data2.xlsx']))
    (HERE / 'cycle_summary.json').write_text(json.dumps(summary, indent=2) + '\n', encoding='utf-8')
    for group, result in summaries.items():
        if result:
            print(group, 'n=', result['n_complete_cycles'], 'median=', round(result['median_h'], 3),
                  'CV=', round(result['observed_cv'], 3), 'AIC=', {k: round(v['aic'], 1) for k, v in result['descriptive_fits'].items()})
    for group, result in classified.items():
        if result and group.startswith('Xiao'):
            print('Classified', group, 'n=', result['n_complete_cycles'],
                  'median=', result['median_h'], 'CV=', round(result['observed_cv'], 3))


if __name__ == '__main__':
    main()
